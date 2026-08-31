"""发现 Shell 在授权工作区内生成或修改的可交付文件。

Office 能力不再通过格式专用模型工具实现。模型按需加载格式 Skill，使用本机
Python/CLI 写文件；这一层只负责在命令结束后做有界差分、格式校验和可信 MIME 判定，
让产物仍能进入 Cowork 的 Artifacts 区。

这里不是文件系统沙箱。Shell 的执行授权与审批仍由 ``run_shell`` 负责；扫描器只读取
已通过 ``filesystem.write`` 授权的 root，而且不跟随符号链接。
"""

from __future__ import annotations

import asyncio
import hashlib
import os
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

import pymupdf
from docx import Document
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pptx import Presentation

from app.cowork.artifact_diff import build_artifact_diff, capture_artifact_baseline
from app.cowork.artifact_formats import TEXT_ARTIFACT_MIME_BY_SUFFIX
from app.cowork.artifact_manifest import legacy_artifact_manifest
from app.cowork.artifact_validation import validate_artifact_in_subprocess

_pymupdf: Any = pymupdf

_SKIPPED_DIRECTORIES = frozenset(
    {
        ".git",
        ".hg",
        ".svn",
        ".workpilot-backups",
        "__pycache__",
        "node_modules",
    }
)
_NATIVE_ARTIFACT_MIME_BY_SUFFIX = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".pdf": "application/pdf",
}
ARTIFACT_MIME_BY_SUFFIX = {
    **TEXT_ARTIFACT_MIME_BY_SUFFIX,
    **_NATIVE_ARTIFACT_MIME_BY_SUFFIX,
}


@dataclass(frozen=True)
class ArtifactFingerprint:
    size_bytes: int
    modified_at_ns: int
    changed_at_ns: int
    baseline_bytes: bytes | None = field(default=None, compare=False, repr=False)


@dataclass(frozen=True)
class WorkspaceArtifactSnapshot:
    files: dict[Path, ArtifactFingerprint]
    truncated: bool = False


@dataclass(frozen=True)
class DiscoveredWorkspaceArtifact:
    path: Path
    title: str
    kind: Literal["file", "report", "table"]
    mime_type: str
    sha256: str
    size_bytes: int
    diff: dict[str, object]
    artifact_manifest: dict[str, object] | None = None


@dataclass(frozen=True)
class WorkspaceArtifactDiscovery:
    artifacts: tuple[DiscoveredWorkspaceArtifact, ...]
    warnings: tuple[str, ...]
    truncated: bool


async def snapshot_workspace_artifacts(
    root: Path,
    *,
    max_scan_entries: int,
    max_files: int,
) -> WorkspaceArtifactSnapshot:
    return await asyncio.to_thread(
        _snapshot_workspace_artifacts,
        root,
        max_scan_entries=max_scan_entries,
        max_files=max_files,
    )


async def discover_workspace_artifacts(
    root: Path,
    *,
    before: WorkspaceArtifactSnapshot,
    max_scan_entries: int,
    max_files: int,
    max_file_bytes: int,
    sandboxed_generated_code: bool = False,
) -> WorkspaceArtifactDiscovery:
    return await asyncio.to_thread(
        _discover_workspace_artifacts,
        root,
        before=before,
        max_scan_entries=max_scan_entries,
        max_files=max_files,
        max_file_bytes=max_file_bytes,
        sandboxed_generated_code=sandboxed_generated_code,
    )


def _snapshot_workspace_artifacts(
    root: Path,
    *,
    max_scan_entries: int,
    max_files: int,
    capture_baselines: bool = True,
) -> WorkspaceArtifactSnapshot:
    canonical_root = root.resolve(strict=True)
    if not canonical_root.is_dir():
        raise ValueError(f"工作区不是现有目录：{canonical_root}")
    candidates: list[tuple[Path, ArtifactFingerprint]] = []
    scanned = 0
    scan_truncated = False
    for current, directories, names in os.walk(canonical_root, followlinks=False):
        directories[:] = sorted(
            name
            for name in directories
            if name not in _SKIPPED_DIRECTORIES and not name.startswith(".")
        )
        for name in sorted(names):
            scanned += 1
            if scanned > max_scan_entries:
                scan_truncated = True
                break
            if name.startswith("."):
                continue
            path = Path(current) / name
            if path.suffix.casefold() not in ARTIFACT_MIME_BY_SUFFIX or path.is_symlink():
                continue
            try:
                resolved = path.resolve(strict=True)
                resolved.relative_to(canonical_root)
                stat = resolved.stat()
            except (OSError, ValueError):
                continue
            if not resolved.is_file():
                continue
            candidates.append(
                (
                    resolved,
                    ArtifactFingerprint(
                        size_bytes=stat.st_size,
                        modified_at_ns=stat.st_mtime_ns,
                        changed_at_ns=stat.st_ctime_ns,
                    ),
                )
            )
        if scan_truncated:
            break

    # 工作区内可能已经积累大量旧交付物。优先保留最近修改的文件，才能在有界快照里
    # 看见本轮新产物，而不是被按目录排序靠前的历史文件挤掉。
    candidates.sort(key=lambda item: (item[1].modified_at_ns, str(item[0])), reverse=True)
    selected = candidates[:max_files]
    if capture_baselines:
        # 单次命令最多为 diff 留 8 MiB 内存，且单文件仍受 2 MiB 上限。扫描再大的
        # 工作区也不会因为“想给右栏看差异”而把 sidecar 内存吃光。
        remaining = 8 * 1024 * 1024
        captured: list[tuple[Path, ArtifactFingerprint]] = []
        for path, fingerprint in selected:
            baseline = capture_artifact_baseline(path, remaining_bytes=remaining)
            if baseline is not None:
                remaining -= len(baseline)
            captured.append(
                (
                    path,
                    ArtifactFingerprint(
                        size_bytes=fingerprint.size_bytes,
                        modified_at_ns=fingerprint.modified_at_ns,
                        changed_at_ns=fingerprint.changed_at_ns,
                        baseline_bytes=baseline,
                    ),
                )
            )
        selected = captured
    files = dict(selected)
    return WorkspaceArtifactSnapshot(
        files=files,
        truncated=scan_truncated or len(candidates) > max_files,
    )


def _discover_workspace_artifacts(
    root: Path,
    *,
    before: WorkspaceArtifactSnapshot,
    max_scan_entries: int,
    max_files: int,
    max_file_bytes: int,
    sandboxed_generated_code: bool = False,
) -> WorkspaceArtifactDiscovery:
    after = _snapshot_workspace_artifacts(
        root,
        max_scan_entries=max_scan_entries,
        max_files=max_files,
        capture_baselines=False,
    )
    changed = sorted(
        path for path, fingerprint in after.files.items() if before.files.get(path) != fingerprint
    )
    artifacts: list[DiscoveredWorkspaceArtifact] = []
    warnings: list[str] = []
    for path in changed:
        try:
            artifacts.append(
                _validated_artifact(
                    path,
                    previous=before.files.get(path),
                    max_file_bytes=max_file_bytes,
                    sandboxed_generated_code=sandboxed_generated_code,
                )
            )
        except (OSError, UnicodeError, ValueError, zipfile.BadZipFile) as error:
            warnings.append(f"{path.name}: {error}")
    return WorkspaceArtifactDiscovery(
        artifacts=tuple(artifacts),
        warnings=tuple(warnings),
        truncated=before.truncated or after.truncated,
    )


def _validated_artifact(
    path: Path,
    *,
    previous: ArtifactFingerprint | None,
    max_file_bytes: int,
    sandboxed_generated_code: bool = False,
) -> DiscoveredWorkspaceArtifact:
    stat = path.stat()
    if stat.st_size > max_file_bytes:
        raise ValueError(f"文件超过产物登记上限 {max_file_bytes} bytes")
    suffix = path.suffix.casefold()
    mime_type = ARTIFACT_MIME_BY_SUFFIX[suffix]
    if suffix in _NATIVE_ARTIFACT_MIME_BY_SUFFIX:
        _validate_native_file(path, suffix, max_file_bytes=max_file_bytes)
    else:
        path.read_bytes().decode("utf-8-sig")
    manifest: dict[str, object] | None = None
    artifact_type = suffix.removeprefix(".")
    if artifact_type in {"docx", "xlsx", "pptx", "pdf", "html"}:
        report = validate_artifact_in_subprocess(
            path,
            render_visual=suffix == ".pptx",
            max_file_bytes=max_file_bytes,
        )
        if not report.deliverable:
            failures = report.quality.warnings[:3]
            raise ValueError("产物安全/结构验证失败：" + "；".join(failures))
        manifest = legacy_artifact_manifest(
            artifact_type=artifact_type,  # type: ignore[arg-type]
            report=report,
            sandboxed_generated_code=sandboxed_generated_code,
        ).model_dump(mode="json")
    return DiscoveredWorkspaceArtifact(
        path=path,
        title=path.name,
        kind=(
            "table" if suffix == ".xlsx" else "report" if suffix in {".docx", ".pdf"} else "file"
        ),
        mime_type=mime_type,
        sha256=_sha256(path),
        size_bytes=stat.st_size,
        diff=build_artifact_diff(
            after_path=path,
            before_bytes=None if previous is None else previous.baseline_bytes,
            created=previous is None,
        ),
        artifact_manifest=manifest,
    )


def _validate_native_file(path: Path, suffix: str, *, max_file_bytes: int) -> None:
    if suffix in {".docx", ".xlsx", ".pptx"}:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10_000:
                raise ValueError("OOXML 压缩包条目过多")
            max_uncompressed_bytes = max(
                max_file_bytes,
                min(max_file_bytes * 10, 512 * 1024 * 1024),
            )
            if sum(item.file_size for item in entries) > max_uncompressed_bytes:
                raise ValueError("OOXML 解压后内容超过验证上限")
            broken = archive.testzip()
        if broken is not None:
            raise ValueError(f"OOXML 压缩包损坏：{broken}")
    if suffix == ".docx":
        Document(str(path))
    elif suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=False)
        workbook.close()
    elif suffix == ".pptx":
        Presentation(str(path))
    else:
        document = _pymupdf.open(path)
        try:
            if document.needs_pass:
                raise ValueError("PDF 已加密，无法验证")
            _ = document.page_count
        finally:
            document.close()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


__all__ = [
    "ARTIFACT_MIME_BY_SUFFIX",
    "DiscoveredWorkspaceArtifact",
    "WorkspaceArtifactDiscovery",
    "WorkspaceArtifactSnapshot",
    "discover_workspace_artifacts",
    "snapshot_workspace_artifacts",
]
